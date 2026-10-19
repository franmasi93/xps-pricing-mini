"""E1-B Perfil comercial e importación de clientes (TE1B-01 a TE1B-24) y 6-B historial de correos y columnas.

Importación con la plantilla oficial (.xlsx y .csv con ; y ,), catálogos del perfil, filtros combinados e interés
calculado, consentimiento en boletines, {{nombre_contacto}} y {{comercial}}, historial de los 3 tipos de correo con los
eventos de Tencent, preferencias de columnas y exportar → importar sin cambios. Correo siempre con el proveedor FALSO.
"""
import base64
import csv
import io
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from pydantic import SecretStr
from sqlalchemy import select

from app.config import get_settings
from app.models import (
    Auditoria,
    CampanaDestinatario,
    CatalogoPerfil,
    Cliente,
    ConfigEmpresa,
    ConfigSistema,
    ContactoCliente,
    Pais,
    EnvioCorreo,
    Notificacion,
    TencentEvento,
)
from app.services import importacion_clientes
from cotizacion_escenarios import Escenario, cargar_condiciones, pais_ecuador
from fabricas import cedula_libre, crear_cliente, crear_contacto, crear_pais
from test_e1_boletines import (  # noqa: F401 - fixtures y ayudas de E1
    TOKEN_EVENTOS,
    _campana,
    _enviar,
    _escenario,
    _hoy,
    _lista,
    _lote,
    _plantilla,
    BOLETIN,
    correo_listo,
)
from test_f4_correo import _con_llaves, _emitida, _sin_red, falso, llave  # noqa: F401 - fixtures de F4
from test_f4_pdf import EMPRESA

OFICIAL = Path(__file__).resolve().parent.parent / "docs" / "extensiones" / "importacion"
VACIAS_PERMITIDAS = {"mas_tarifas"}  # además de las filas de tarifas sin uso (fila_NN)


def _b64(contenido: bytes) -> str:
    return base64.b64encode(contenido).decode()


def _codigo(respuesta) -> str:
    return respuesta.json()["detail"]["codigo"]


def _fila_cliente(comercial_email: str, identificacion: str, **cambios) -> dict[str, str]:
    fila = {c: "" for c in importacion_clientes.COLUMNAS_CLIENTE}
    fila |= {"tipo_identificacion": "CEDULA", "identificacion": identificacion, "razon_social": f"PRUEBA {identificacion}",
             "pais": "EC", "ciudad": "Guayaquil", "comercial_email": comercial_email, "tipo_cliente": "IMPORTADOR",
             "sector": "TEXTIL", "servicios_interes": "FCL|LCL", "rutas_interes": "CN>EC", "acepta_promociones": "SI",
             "fecha_consentimiento": "19/10/2026", "medio_consentimiento": "FORMULARIO_WEB"}
    return fila | cambios


def _fila_contacto(identificacion: str, email: str, **cambios) -> dict[str, str]:
    fila = {c: "" for c in importacion_clientes.COLUMNAS_CONTACTO}
    fila |= {"identificacion_cliente": identificacion, "nombre": f"Contacto {email.split('@')[0]}", "email": email,
             "area": "COMPRAS"}
    return fila | cambios


def _xlsx(clientes: list[dict], contactos: list[dict], *, con_ejemplos: bool = False) -> bytes:
    """Sobre la plantilla oficial: encabezado y fila de descripción; las de ejemplo se dejan o se borran."""
    libro = load_workbook(OFICIAL / "plantilla_importacion_clientes.xlsx")
    for nombre, filas, columnas in (("Clientes", clientes, importacion_clientes.COLUMNAS_CLIENTE),
                                    ("Contactos", contactos, importacion_clientes.COLUMNAS_CONTACTO)):
        hoja = libro[nombre]
        inicio = 5 if con_ejemplos else 3
        if not con_ejemplos:
            for fila in hoja.iter_rows(min_row=3, max_row=hoja.max_row):
                for celda in fila:
                    celda.value = None
        for numero, datos in enumerate(filas, start=inicio):
            for indice, columna in enumerate(columnas, start=1):
                hoja.cell(row=numero, column=indice, value=datos[columna] or None)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _csv(clientes: list[dict], contactos: list[dict], separador: str = ";", bom: bool = True) -> bytes:
    por_id = {c["identificacion"]: c for c in clientes}
    salida = io.StringIO()
    escritor = csv.writer(salida, delimiter=separador, lineterminator="\n")
    escritor.writerow(importacion_clientes.COLUMNAS_CSV)
    for contacto in contactos:
        cliente = por_id[contacto["identificacion_cliente"]]
        escritor.writerow([cliente[c] for c in importacion_clientes.COLUMNAS_CLIENTE]
                          + [contacto[c] for c in importacion_clientes.COLUMNAS_CONTACTO[1:]])
    return (("﻿" if bom else "") + salida.getvalue()).encode("utf-8")


def _previa(client, h, nombre: str, contenido: bytes):
    return client.post("/v1/clientes/importacion/vista-previa", json={"nombre": nombre, "contenido_base64": _b64(contenido)},
                       headers=h)


def _confirmar(client, h, nombre: str, contenido: bytes, modo: str = "CREAR_Y_ACTUALIZAR"):
    return client.post("/v1/clientes/importacion/confirmar",
                       json={"nombre": nombre, "contenido_base64": _b64(contenido), "modo": modo}, headers=h)


def _cliente(db_session, identificacion: str) -> Cliente:
    db_session.expire_all()
    return db_session.scalars(select(Cliente).where(Cliente.ruc_o_id == identificacion)).one()


@pytest.fixture
def base(db_session, crear_usuario, headers_de):
    """Ecuador y China activos (las rutas CN>EC), un comercial con correo propio y el administrador que importa."""
    pais_ecuador(db_session)
    if db_session.scalars(select(Pais).where(Pais.iso2 == "CN")).one_or_none() is None:
        crear_pais(db_session, iso2="CN")
    comercial = crear_usuario("comercial", nombre="Carla Comercial")
    admin = crear_usuario("admin")
    return {"comercial": comercial, "admin": admin, "h": headers_de(admin)}


# ---------- Migración, catálogos y permisos (TE1B-01 a TE1B-03) ----------


def test_clientes_existentes_quedan_con_consentimiento_y_catalogos_sembrados(db_session, crear_usuario):  # TE1B-01
    cliente = crear_cliente(db_session, crear_usuario("comercial"))
    contacto = crear_contacto(db_session, cliente)
    db_session.refresh(cliente)
    db_session.refresh(contacto)
    assert (cliente.estado_comercial, cliente.acepta_promociones, cliente.servicios_interes) == ("ACTIVO", True, [])
    assert (contacto.recibe_promociones, contacto.recibir_avisos) == (True, True)
    tipos = set(db_session.scalars(select(CatalogoPerfil.tipo)))
    assert {"TIPO_CLIENTE", "ESTADO_COMERCIAL", "SECTOR", "FRECUENCIA", "VOLUMEN", "ORIGEN_CLIENTE", "AREA_CONTACTO",
            "MEDIO_CONSENTIMIENTO", "TIPO_CARGA"} <= tipos
    medio = db_session.scalars(select(CatalogoPerfil).where(CatalogoPerfil.codigo == "RELACION_COMERCIAL_PREVIA")).one()
    assert medio.tipo == "MEDIO_CONSENTIMIENTO"
    assert all(e.fijo for e in db_session.scalars(select(CatalogoPerfil).where(CatalogoPerfil.tipo == "ESTADO_COMERCIAL")))
    assert db_session.get(ConfigSistema, 1).importacion_max_filas == 2000


@pytest.mark.parametrize(("rol", "plantilla", "catalogo"), [
    ("admin", 200, 201), ("comercial", 200, 403), ("pricing", 403, 403), ("gerencia", 403, 403),
])
def test_permisos_de_importacion_y_catalogos(client, crear_usuario, headers_de, rol, plantilla, catalogo):  # TE1B-02
    h = headers_de(crear_usuario(rol))
    assert client.get("/v1/clientes/importacion/plantilla.xlsx", headers=h).status_code == plantilla
    assert client.post("/v1/clientes/importacion/vista-previa", json={"nombre": "x.csv", "contenido_base64": ""},
                       headers=h).status_code in ((403,) if plantilla == 403 else (422,))
    nuevo = client.post("/v1/catalogos-perfil", json={"tipo": "SECTOR", "nombre": f"Sector {rol}"}, headers=h)
    assert nuevo.status_code == catalogo
    assert client.get("/v1/catalogos-perfil?tipo=SECTOR", headers=h).status_code == 200


def test_catalogo_crud_con_auditoria_y_estados_fijos(client, base, db_session):  # TE1B-03
    h = base["h"]
    nuevo = client.post("/v1/catalogos-perfil", json={"tipo": "SECTOR", "nombre": "Minería y petróleo"}, headers=h)
    assert nuevo.status_code == 201, nuevo.text
    assert nuevo.json()["codigo"] == "MINERIA_Y_PETROLEO"
    editado = client.patch(f"/v1/catalogos-perfil/{nuevo.json()['id']}", json={"nombre": "Minería", "activo": False}, headers=h)
    assert editado.status_code == 200 and editado.json()["activo"] is False
    assert db_session.scalars(select(Auditoria).where(Auditoria.entidad == "catalogo_perfil",
                                                      Auditoria.entidad_id == str(nuevo.json()["id"]))).all()
    fijo = db_session.scalars(select(CatalogoPerfil).where(CatalogoPerfil.tipo == "ESTADO_COMERCIAL")).first()
    assert client.patch(f"/v1/catalogos-perfil/{fijo.id}", json={"nombre": "Otro"}, headers=h).status_code == 409
    assert client.post("/v1/catalogos-perfil", json={"tipo": "ESTADO_COMERCIAL", "nombre": "Nuevo"}, headers=h).status_code in (409, 422)


# ---------- Plantilla oficial (TE1B-04) ----------


def test_plantilla_descargada_coincide_con_la_oficial(client, base, db_session):  # TE1B-04
    respuesta = client.get("/v1/clientes/importacion/plantilla.xlsx", headers=base["h"])
    assert respuesta.status_code == 200
    descargada = load_workbook(io.BytesIO(respuesta.content))
    oficial = load_workbook(OFICIAL / "plantilla_importacion_clientes.xlsx")
    assert descargada.sheetnames == oficial.sheetnames
    for hoja in ("Clientes", "Contactos"):
        assert [c.value for c in descargada[hoja][1]] == [c.value for c in oficial[hoja][1]]
    encabezados = lambda libro, hoja: [importacion_clientes._encabezado(c.value) for c in libro[hoja][1] if c.value]  # noqa: E731
    assert encabezados(descargada, "Clientes") == list(importacion_clientes.COLUMNAS_CLIENTE)
    assert encabezados(descargada, "Contactos") == list(importacion_clientes.COLUMNAS_CONTACTO)
    catalogos = {str(c[0].value).split(" ")[0]: [v.value for v in c[1:] if v.value] for c in descargada["Catalogos"].iter_cols()}
    assert "IMPORTADOR" in catalogos["tipo_cliente"] and "RELACION_COMERCIAL_PREVIA" in catalogos["medio_consentimiento"]
    assert "CEDULA" in catalogos["tipo_identificacion"] and "FCL" in catalogos["servicios"]
    assert descargada["Clientes"].data_validations.dataValidation  # las listas desplegables siguen ahí
    csv_oficial = (OFICIAL / "plantilla_importacion_clientes.csv").read_text(encoding="utf-8-sig").splitlines()[0]
    csv_descargado = client.get("/v1/clientes/importacion/plantilla.csv", headers=base["h"]).content.decode("utf-8-sig")
    assert csv_descargado.splitlines()[0].split(";") == csv_oficial.split(";") == list(importacion_clientes.COLUMNAS_CSV)


def test_la_plantilla_oficial_viaja_en_la_imagen():  # TE1B-04b
    raiz = OFICIAL.parent.parent.parent
    assert "COPY docs/extensiones/importacion ./docs/extensiones/importacion" in (raiz / "Dockerfile").read_text(encoding="utf-8")
    assert "!docs/extensiones/importacion/" in (raiz / ".gcloudignore").read_text(encoding="utf-8")
    assert "!docs/extensiones/importacion/" in (raiz / ".dockerignore").read_text(encoding="utf-8")
    assert importacion_clientes.PLANTILLA == OFICIAL / "plantilla_importacion_clientes.xlsx"
    # --set-secrets reemplaza la lista: el comando de la guía debe llevar los 5 (el 2026-10-19 faltaron 2 y los jobs dieron 503).
    comando = next(l for l in (raiz / "docs" / "DESPLIEGUE.md").read_text(encoding="utf-8").splitlines() if "gcloud run deploy" in l)
    for secreto in ("DATABASE_URL=", "JWT_SECRET=", "CONFIG_CIFRADO_KEY=", "TAREAS_SECRETO=", "TENCENT_EVENTOS_TOKEN="):
        assert secreto in comando


# ---------- Importación (TE1B-05 a TE1B-11) ----------


@pytest.mark.parametrize("formato", ["xlsx", "csv;", "csv,", "csv-sin-bom"])
def test_importa_xlsx_y_csv(client, base, db_session, formato):  # TE1B-05
    email = base["comercial"].email
    a, b = cedula_libre(db_session), cedula_libre(db_session)
    clientes = [_fila_cliente(email, a, etiquetas="vip|feria"), _fila_cliente(email, b, tipo_cliente="EXPORTADOR",
                                                                            acepta_promociones="No")]
    contactos = [_fila_contacto(a, f"uno.{a}@prueba.ec", principal="SI"), _fila_contacto(a, f"dos.{a}@prueba.ec",
                                                                                         recibe_promociones="NO"),
                 _fila_contacto(b, f"tres.{b}@prueba.ec", recibe_cotizaciones="0")]
    if formato == "xlsx":
        nombre, contenido = "clientes.xlsx", _xlsx(clientes, contactos)
    else:
        nombre = "clientes.csv"
        contenido = _csv(clientes, contactos, ";" if formato != "csv," else ",", bom=formato != "csv-sin-bom")
    previa = _previa(client, base["h"], nombre, contenido)
    assert previa.status_code == 200, previa.text
    resumen = previa.json()["resumen"]
    assert (resumen["clientes_nuevos"], resumen["contactos_nuevos"], resumen["clientes_error"], resumen["contactos_error"]) == (2, 3, 0, 0)
    assert db_session.scalars(select(Cliente).where(Cliente.ruc_o_id == a)).one_or_none() is None  # la vista previa no guarda
    confirmado = _confirmar(client, base["h"], nombre, contenido)
    assert confirmado.status_code == 200, confirmado.text
    assert (confirmado.json()["resumen"]["clientes_creados"], confirmado.json()["resumen"]["contactos_creados"]) == (2, 3)
    primero, segundo = _cliente(db_session, a), _cliente(db_session, b)
    assert (primero.tipo_cliente, primero.sector, primero.servicios_interes, primero.rutas_interes) == (
        "IMPORTADOR", "TEXTIL", ["FCL", "LCL"], ["CN>EC"])
    assert primero.etiquetas == ["VIP", "FERIA"] and primero.estado_comercial == "ACTIVO"
    assert primero.comercial_asignado_id == base["comercial"].id and primero.fecha_consentimiento.isoformat() == "2026-10-19"
    assert segundo.acepta_promociones is False and segundo.tipo_cliente == "EXPORTADOR"
    por_email = {c.email: c for c in primero.contactos + segundo.contactos}
    assert por_email[f"uno.{a}@prueba.ec"].principal is True and por_email[f"uno.{a}@prueba.ec"].area == "COMPRAS"
    assert por_email[f"dos.{a}@prueba.ec"].recibe_promociones is False
    assert por_email[f"tres.{b}@prueba.ec"].recibir_avisos is False and por_email[f"tres.{b}@prueba.ec"].recibe_promociones is True
    auditoria = db_session.scalars(select(Auditoria).where(Auditoria.entidad == "importacion_clientes")).all()
    assert auditoria and auditoria[-1].valores_despues["clientes_creados"] == 2


def test_errores_por_fila_y_columna(client, base, db_session, crear_usuario):  # TE1B-06
    email = base["comercial"].email
    inactivo = crear_usuario("comercial", activo=False)
    ids = [cedula_libre(db_session) for _ in range(6)]
    clientes = [
        _fila_cliente(email, ids[0], pais="ZZ"),
        _fila_cliente(inactivo.email, ids[1]),
        _fila_cliente(email, ids[2], tipo_cliente="ASTRONAUTA", acepta_promociones="quizá"),
        _fila_cliente(email, "123", tipo_identificacion="CEDULA"),
        _fila_cliente(email, ids[4], razon_social=""),
        _fila_cliente(email, ids[5]),
        _fila_cliente(email, ids[5], razon_social="REPETIDO"),
    ]
    contactos = [_fila_contacto(ids[5], "no-es-correo"), _fila_contacto(ids[5], f"a.{ids[5]}@prueba.ec", principal="SI"),
                 _fila_contacto(ids[5], f"b.{ids[5]}@prueba.ec", principal="SI"),
                 _fila_contacto(ids[5], f"a.{ids[5]}@prueba.ec"), _fila_contacto("9999999999", "x@prueba.ec")]
    previa = _previa(client, base["h"], "errores.xlsx", _xlsx(clientes, contactos)).json()
    errores = {(f["fila"], f["tipo"]): {e["columna"] for e in f["errores"]} for f in previa["filas"]}
    assert errores[(3, "cliente")] == {"pais"}
    assert errores[(4, "cliente")] == {"comercial_email"}
    assert errores[(5, "cliente")] == {"tipo_cliente", "acepta_promociones"}
    assert errores[(6, "cliente")] == {"identificacion"}
    assert errores[(7, "cliente")] == {"razon_social"}
    assert errores[(8, "cliente")] == set() and errores[(9, "cliente")] == {"identificacion"}  # repetida en el archivo
    assert errores[(3, "contacto")] == {"email"}
    assert errores[(4, "contacto")] == set() and errores[(5, "contacto")] == {"principal"}
    assert errores[(6, "contacto")] == {"email"}  # repetido para el mismo cliente
    assert errores[(7, "contacto")] == {"identificacion_cliente"}
    mensaje = next(f for f in previa["filas"] if (f["fila"], f["tipo"]) == (3, "cliente"))["errores"][0]["mensaje"]
    assert "ZZ" in mensaje and "ISO" in mensaje
    confirmado = _confirmar(client, base["h"], "errores.xlsx", _xlsx(clientes, contactos)).json()["resumen"]
    assert confirmado["clientes_creados"] == 1 and confirmado["errores"] == 10  # las filas buenas entran, las malas no


def test_duplicado_contra_la_base_actualiza_o_se_omite(client, base, db_session):  # TE1B-07
    comercial = base["comercial"]
    existente = crear_cliente(db_session, comercial, razon_social="CLIENTE VIEJO S.A.")
    existente.ciudad, existente.telefono = "Quito", "022000000"
    contacto = crear_contacto(db_session, existente, nombre="Contacto Viejo")
    contacto.email = f"viejo.{existente.id}@prueba.ec"
    db_session.commit()
    fila = _fila_cliente(comercial.email, existente.ruc_o_id, razon_social="CLIENTE VIEJO S.A.", ciudad="Cuenca", telefono="")
    contactos = [_fila_contacto(existente.ruc_o_id, contacto.email, nombre="Contacto Renombrado", area="")]
    archivo = _xlsx([fila], contactos)
    previa = _previa(client, base["h"], "dup.xlsx", archivo).json()
    acciones = {f["tipo"]: (f["accion"], f["cambios"]) for f in previa["filas"]}
    assert acciones["cliente"][0] == "ACTUALIZAR" and "ciudad" in acciones["cliente"][1] and "telefono" not in acciones["cliente"][1]
    assert acciones["contacto"] == ("ACTUALIZAR", ["nombre"])
    solo_nuevos = _confirmar(client, base["h"], "dup.xlsx", archivo, "SOLO_NUEVOS").json()["resumen"]
    assert (solo_nuevos["clientes_actualizados"], solo_nuevos["omitidos"]) == (0, 2)
    assert _cliente(db_session, existente.ruc_o_id).ciudad == "Quito"
    actualizado = _confirmar(client, base["h"], "dup.xlsx", archivo, "CREAR_Y_ACTUALIZAR").json()["resumen"]
    assert (actualizado["clientes_actualizados"], actualizado["contactos_actualizados"]) == (1, 1)
    cliente = _cliente(db_session, existente.ruc_o_id)
    assert (cliente.ciudad, cliente.telefono, cliente.tipo_cliente) == ("Cuenca", "022000000", "IMPORTADOR")  # vacío no borra
    assert cliente.contactos[0].nombre == "Contacto Renombrado" and len(cliente.contactos) == 1  # nunca duplica
    assert _previa(client, base["h"], "dup.xlsx", archivo).json()["resumen"]["cambios"] == 0


def test_filas_de_ejemplo_y_descripcion_se_ignoran(client, base, db_session):  # TE1B-08
    a = cedula_libre(db_session)
    contenido = _xlsx([_fila_cliente(base["comercial"].email, a)], [_fila_contacto(a, f"c.{a}@prueba.ec")], con_ejemplos=True)
    resumen = _previa(client, base["h"], "con_ejemplos.xlsx", contenido).json()["resumen"]
    assert (resumen["ejemplos_ignorados"], resumen["clientes_nuevos"], resumen["contactos_nuevos"], resumen["clientes_error"]) == (4, 1, 1, 0)
    ejemplo_csv = client.get("/v1/clientes/importacion/plantilla.csv", headers=base["h"]).content
    solo_ejemplos = _previa(client, base["h"], "ejemplo.csv", ejemplo_csv).json()["resumen"]
    assert (solo_ejemplos["ejemplos_ignorados"], solo_ejemplos["cambios"]) == (4, 0)


def test_limite_de_filas_configurable(client, base, db_session):  # TE1B-09
    db_session.get(ConfigSistema, 1).importacion_max_filas = 2
    db_session.commit()
    ids = [cedula_libre(db_session) for _ in range(3)]
    contenido = _xlsx([_fila_cliente(base["comercial"].email, i) for i in ids], [])
    assert _codigo(_previa(client, base["h"], "grande.xlsx", contenido)) == "DEMASIADAS_FILAS"
    assert _previa(client, base["h"], "chico.xlsx", _xlsx([_fila_cliente(base["comercial"].email, ids[0])], [])).status_code == 200


def test_archivo_invalido_y_columnas_faltantes(client, base):  # TE1B-10
    assert _codigo(_previa(client, base["h"], "notas.txt", b"hola")) == "ARCHIVO_INVALIDO"
    assert _codigo(_previa(client, base["h"], "malo.xlsx", b"no es excel")) == "ARCHIVO_INVALIDO"
    respuesta = _previa(client, base["h"], "corto.csv", "identificacion;razon_social\n1;X\n".encode())
    assert _codigo(respuesta) == "COLUMNAS_FALTANTES" and "tipo_identificacion" in respuesta.json()["detail"]["mensaje"]


def test_comercial_solo_importa_sus_clientes_y_a_lista(client, base, db_session, crear_usuario, headers_de):  # TE1B-11
    comercial = base["comercial"]
    h = headers_de(comercial)
    otro = crear_usuario("comercial")
    a, b = cedula_libre(db_session), cedula_libre(db_session)
    previa = _previa(client, h, "c.xlsx", _xlsx([_fila_cliente(comercial.email, a), _fila_cliente(otro.email, b)], [])).json()
    assert [f["accion"] for f in previa["filas"]] == ["CREAR", "ERROR"]
    pricing = crear_usuario("pricing")
    lista = _lista(client, headers_de(pricing), [], nombre="Lista importada")
    contenido = _xlsx([_fila_cliente(comercial.email, a)], [_fila_contacto(a, f"l.{a}@prueba.ec")])
    cuerpo = {"nombre": "l.xlsx", "contenido_base64": _b64(contenido), "modo": "CREAR_Y_ACTUALIZAR"}
    assert client.post(f"/v1/boletines/listas/{lista['id']}/importar-archivo", json=cuerpo,
                       headers=headers_de(pricing)).status_code == 403  # pricing no edita clientes
    admin = base["h"]
    importado = client.post(f"/v1/boletines/listas/{lista['id']}/importar-archivo", json=cuerpo, headers=admin)
    assert importado.status_code == 200, importado.text
    assert importado.json()["resumen"]["agregados_a_lista"] == 1
    detalle = client.get(f"/v1/boletines/listas/{lista['id']}", headers=admin).json()
    assert [m["email"] for m in detalle["resumen"]["muestra"]] == [f"l.{a}@prueba.ec"]
    otra_vez = client.post(f"/v1/boletines/listas/{lista['id']}/importar-archivo", json=cuerpo, headers=admin).json()
    assert otra_vez["resumen"]["agregados_a_lista"] == 0  # ya estaba: no se duplica


# ---------- Ficha, filtros e interés calculado (TE1B-12 a TE1B-15) ----------


def test_ficha_guarda_el_perfil_y_valida_catalogos(client, base, db_session):  # TE1B-12
    pais = pais_ecuador(db_session)
    cuerpo = {"razon_social": "PERFIL S.A.", "tipo_id": "CEDULA", "ruc_o_id": cedula_libre(db_session), "pais_id": pais.id,
              "comercial_asignado_id": base["comercial"].id, "tipo_cliente": "IMPORTADOR", "sector": "TEXTIL",
              "servicios_interes": ["FCL"], "rutas_interes": ["CN>EC"], "etiquetas": ["vip"], "acepta_promociones": True}
    creado = client.post("/v1/clientes", json=cuerpo, headers=base["h"])
    assert creado.status_code == 201, creado.text
    assert (creado.json()["estado_comercial"], creado.json()["etiquetas"]) == ("ACTIVO", ["VIP"])
    malo = client.patch(f"/v1/clientes/{creado.json()['id']}", json={"sector": "INVENTADO"}, headers=base["h"])
    assert malo.status_code == 422
    contacto = client.post(f"/v1/clientes/{creado.json()['id']}/contactos", json={
        "nombre": "Ana", "email": "ana@perfil.ec", "area": "COMPRAS", "recibe_promociones": False, "recibe_cotizaciones": False},
        headers=base["h"])
    assert contacto.status_code == 201, contacto.text
    assert (contacto.json()["area"], contacto.json()["recibe_promociones"], contacto.json()["recibe_cotizaciones"]) == (
        "COMPRAS", False, False)


def test_filtros_combinados_de_la_lista_de_clientes(client, base, db_session):  # TE1B-13
    comercial = base["comercial"]
    datos = [("IMPORTADOR", "TEXTIL", ["FCL"], True), ("IMPORTADOR", "ALIMENTOS", ["AEREO"], True),
             ("EXPORTADOR", "TEXTIL", ["FCL", "LCL"], False)]
    clientes = []
    for tipo, sector, servicios, acepta in datos:
        cliente = crear_cliente(db_session, comercial)
        cliente.tipo_cliente, cliente.sector, cliente.servicios_interes, cliente.acepta_promociones = tipo, sector, servicios, acepta
        cliente.etiquetas = ["VIP"] if sector == "TEXTIL" else []
        clientes.append(cliente)
    db_session.commit()
    ids = lambda consulta: {c["id"] for c in client.get(f"/v1/clientes?{consulta}", headers=base["h"]).json()} & {c.id for c in clientes}  # noqa: E731
    assert ids("tipo=IMPORTADOR") == {clientes[0].id, clientes[1].id}
    assert ids("tipo=IMPORTADOR&sector=TEXTIL") == {clientes[0].id}
    assert ids("servicio=FCL") == {clientes[0].id, clientes[2].id}
    assert ids("servicio=LCL|AEREO") == {clientes[1].id, clientes[2].id}
    assert ids("etiqueta=vip&acepta_promociones=NO") == {clientes[2].id}


def test_interes_calculado_y_listas_por_filtro_y_o(client, base, db_session, crear_usuario, headers_de):  # TE1B-14
    comercial = crear_usuario("comercial")
    esc = Escenario(db_session, comercial, modo="FCL")
    esc.cliente.tipo_cliente = "IMPORTADOR"
    esc.contacto.email = f"fcl.{esc.cliente.id}@prueba.ec"
    otro = crear_cliente(db_session, comercial)
    otro.tipo_cliente, otro.sector = "EXPORTADOR", "TEXTIL"
    contacto_otro = crear_contacto(db_session, otro)
    contacto_otro.email = f"exp.{otro.id}@prueba.ec"
    db_session.commit()
    cargar_condiciones(db_session)
    hc = headers_de(comercial)
    borrador = client.post("/v1/cotizaciones", json=esc.payload("FOB"), headers=hc)
    assert borrador.status_code == 201, borrador.text
    assert client.post(f"/v1/cotizaciones/{borrador.json()['id']}/proponer-lineas", json={}, headers=hc).status_code == 200
    assert client.post(f"/v1/cotizaciones/{borrador.json()['id']}/emitir", json={}, headers=hc).status_code == 200
    interes = client.get(f"/v1/clientes/{esc.cliente.id}/interes", headers=base["h"]).json()
    assert [s["codigo"] for s in interes["servicios"]] == ["FCL"] and interes["servicios"][0]["veces"] == 1
    assert len(interes["rutas"]) == 1 and ">" in interes["rutas"][0]["codigo"]
    h = headers_de(crear_usuario("pricing"))

    def emails(condiciones, combinar="Y"):
        respuesta = client.post("/v1/boletines/listas/vista-previa", json={
            "tipo": "FILTRO", "filtros": {"condiciones": condiciones, "combinar": combinar}}, headers=h)
        assert respuesta.status_code == 200, respuesta.text
        return {m["email"] for m in respuesta.json()["muestra"]}

    fcl = {"campo": "cotizo_servicio", "valores": ["FCL"], "meses": 6}
    assert esc.contacto.email in emails([fcl]) and contacto_otro.email not in emails([fcl])
    assert emails([fcl, {"campo": "tipo_cliente", "valores": ["EXPORTADOR"]}]) & {esc.contacto.email, contacto_otro.email} == set()
    assert emails([fcl, {"campo": "tipo_cliente", "valores": ["EXPORTADOR"]}], "O") >= {esc.contacto.email, contacto_otro.email}
    assert esc.contacto.email not in emails([{"campo": "cotizo_servicio", "valores": ["AEREO"], "meses": 6}])
    ruta = interes["rutas"][0]["codigo"]
    assert esc.contacto.email in emails([{"campo": "cotizo_ruta", "valores": [ruta], "meses": 6}])
    lista = client.post("/v1/boletines/listas", json={"nombre": "Importadores FCL", "tipo": "FILTRO", "filtros": {
        "condiciones": [fcl, {"campo": "tipo_cliente", "valores": ["IMPORTADOR"]}], "combinar": "Y"}}, headers=h)
    assert lista.status_code == 201, lista.text
    assert lista.json()["filtros"]["combinar"] == "Y"


def test_filtro_por_area_y_valores_desconocidos(client, crear_usuario, headers_de, db_session):  # TE1B-15
    h = headers_de(crear_usuario("pricing"))
    malo = client.post("/v1/boletines/listas/vista-previa", json={"tipo": "FILTRO", "filtros": {
        "condiciones": [{"campo": "color", "valores": ["rojo"]}]}}, headers=h)
    assert _codigo(malo) == "FILTRO_INVALIDO"
    cliente = crear_cliente(db_session, crear_usuario("comercial"))
    contacto = crear_contacto(db_session, cliente)
    contacto.email, contacto.area = f"log.{cliente.id}@prueba.ec", "LOGISTICA"
    db_session.commit()
    respuesta = client.post("/v1/boletines/listas/vista-previa", json={"tipo": "FILTRO", "filtros": {
        "condiciones": [{"campo": "area_contacto", "valores": ["LOGISTICA"]}]}}, headers=h).json()
    assert contacto.email in {m["email"] for m in respuesta["muestra"]}


# ---------- Consentimiento y variables de Tencent (TE1B-16 a TE1B-18) ----------


def _contacto_con_email(db_session, cliente, email, nombre, **campos):
    contacto = crear_contacto(db_session, cliente, nombre=nombre)
    contacto.email = email
    for campo, valor in campos.items():
        setattr(contacto, campo, valor)
    db_session.commit()
    return contacto


def test_boletin_respeta_el_consentimiento(client, crear_usuario, headers_de, db_session, correo_listo):  # TE1B-16
    pricing = crear_usuario("pricing", nombre="Pricing Envía")
    h = headers_de(pricing)
    _, tarifa = _escenario(db_session, crear_usuario("comercial"))
    acepta = crear_cliente(db_session, pricing)
    no_acepta = crear_cliente(db_session, pricing)
    no_acepta.acepta_promociones = False
    db_session.commit()
    si = _contacto_con_email(db_session, acepta, f"si.{acepta.id}@prueba.ec", "Sí Recibe")
    no_contacto = _contacto_con_email(db_session, acepta, f"no.{acepta.id}@prueba.ec", "No Recibe", recibe_promociones=False)
    no_cliente = _contacto_con_email(db_session, no_acepta, f"nc.{no_acepta.id}@prueba.ec", "Cliente No")
    lista = _lista(client, h, [si, no_contacto, no_cliente])
    campana = _campana(client, h, db_session, [tarifa], [lista])
    conteo = client.get(f"/v1/boletines/campanas/{campana['id']}/conteo", headers=h).json()
    assert (conteo["omitidos_consentimiento"], conteo["recibiran"]) == (2, 1)
    assert _enviar(client, h, campana["id"], 1).status_code == 200
    _lote(db_session)
    assert [m.destinatario for m in correo_listo.enviados] == [si.email]
    previa = client.post("/v1/boletines/listas/vista-previa", json={"tipo": "MANUAL", "contactos": [si.id, no_contacto.id, no_cliente.id]},
                         headers=h).json()
    assert previa["sin_consentimiento"] == 2


def test_nombre_contacto_y_comercial_asignado_sin_variables_vacias(client, crear_usuario, headers_de, db_session,
                                                                    correo_listo):  # TE1B-17
    fila = db_session.get(ConfigEmpresa, 1)
    for campo, valor in EMPRESA.items():
        setattr(fila, campo, valor)
    db_session.commit()
    pricing = crear_usuario("pricing", nombre="Pricing Envía")
    h = headers_de(pricing)
    _, tarifa = _escenario(db_session, crear_usuario("comercial"))
    carla = crear_usuario("comercial", nombre="Carla Asignada")
    cliente = crear_cliente(db_session, carla)
    ana = _contacto_con_email(db_session, cliente, f"ana.{cliente.id}@prueba.ec", "Ana Torres")
    campana = _campana(client, h, db_session, [tarifa], [_lista(client, h, [ana])])
    previa = client.get(f"/v1/boletines/campanas/{campana['id']}/vista-previa", headers=h).json()
    assert previa["como"] == {"nombre_contacto": "Ana Torres", "comercial": "Carla Asignada"}  # la pantalla = lo enviado
    assert (previa["variables"]["nombre_contacto"], previa["variables"]["comercial"]) == ("Ana Torres", "Carla Asignada")
    prueba = client.post(f"/v1/boletines/campanas/{campana['id']}/prueba", json={}, headers=h)
    assert prueba.status_code == 200, prueba.text
    assert prueba.json()["como"] == {"nombre_contacto": "Ana Torres", "comercial": "Carla Asignada", "empresa": cliente.razon_social}
    mensaje = correo_listo.enviados[-1]
    assert mensaje.destinatario == pricing.email  # la prueba va a quien la pide, nunca al cliente
    assert (mensaje.datos_plantilla["nombre_contacto"], mensaje.datos_plantilla["comercial"]) == ("Ana Torres", "Carla Asignada")
    _enviar(client, h, campana["id"], 1)
    _lote(db_session)
    real = correo_listo.enviados[-1]
    assert real.destinatario == ana.email
    assert (real.datos_plantilla["nombre_contacto"], real.datos_plantilla["comercial"]) == ("Ana Torres", "Carla Asignada")
    for datos in (mensaje.datos_plantilla, real.datos_plantilla):
        vacias = {k for k, v in datos.items() if v in ("", None) and not k.startswith("fila_")} - VACIAS_PERMITIDAS
        assert vacias == set()
    assert set(real.datos_plantilla) == set(_plantilla(db_session, BOLETIN).variables_esperadas)


def test_sin_comercial_asignado_usa_quien_envia(db_session, crear_usuario):  # TE1B-18
    from app.services import campanas

    pricing = crear_usuario("pricing", nombre="Pricing Envía")
    cliente = crear_cliente(db_session, crear_usuario("comercial", nombre="Comercial Cliente"))
    contacto = _contacto_con_email(db_session, cliente, f"x.{cliente.id}@prueba.ec", "X")
    assert campanas.comercial_del_contacto(db_session, contacto.id, pricing.nombre) == "Comercial Cliente"
    assert campanas.comercial_del_contacto(db_session, None, pricing.nombre) == "Pricing Envía"


# ---------- 6-B Historial de correos, eventos y columnas (TE1B-19 a TE1B-23) ----------


def test_historial_de_los_tres_tipos_y_eventos(client, crear_usuario, headers_de, db_session, correo_listo,
                                               monkeypatch):  # TE1B-19
    monkeypatch.setattr(get_settings(), "TENCENT_EVENTOS_TOKEN", SecretStr(TOKEN_EVENTOS))
    pricing = crear_usuario("pricing", nombre="Pricing Envía")
    h = headers_de(pricing)
    comercial = crear_usuario("comercial")
    _, tarifa = _escenario(db_session, crear_usuario("comercial"))
    emitida, esc = _emitida(client, headers_de(comercial), db_session, comercial, "hist@prueba.ec")
    cliente, contacto = esc.cliente, esc.contacto
    db_session.add(EnvioCorreo(cotizacion_id=emitida["id"], destinatario=contacto.email, cc=[], asunto="Cotización X",
                               adjunto_nombre="x.pdf", adjunto_bytes=1, estado="ENVIADO", message_id="msg-hist-cot",
                               enviado_por=comercial.id))
    db_session.add(Notificacion(evento="COTIZACION_POR_VENCER", audiencia="EXTERNA", entidad_tipo="cotizacion",
                                entidad_id=str(emitida["id"]), cliente_id=cliente.id, email=contacto.email,
                                canal="CORREO", estado="ENVIADO", titulo="Su cotización vence", asunto="Aviso de vencimiento",
                                clave_idempotencia=f"hist-{cliente.id}", message_id="msg-hist-aviso",
                                enviado_en=datetime.now(timezone.utc) - timedelta(hours=1)))
    db_session.commit()
    campana = _campana(client, h, db_session, [tarifa], [_lista(client, h, [contacto])])
    _enviar(client, h, campana["id"], 1)
    _lote(db_session)
    destinatario = db_session.scalars(select(CampanaDestinatario).where(CampanaDestinatario.campana_id == campana["id"])).one()
    ruta = f"/v1/clientes/{cliente.id}/correos"
    historial = client.get(ruta, headers=headers_de(comercial)).json()
    assert {i["tipo"] for i in historial} == {"COTIZACION", "BOLETIN", "AVISO"}
    assert all(i["contacto"] == contacto.nombre and i["estado_entrega"] == "SIN_EVENTOS" for i in historial)
    assert {i["tipo"] for i in client.get(f"{ruta}?tipo=BOLETIN", headers=h).json()} == {"BOLETIN"}
    for evento in ({"event": "delivery", "bulkId": destinatario.message_id, "email": contacto.email},
                   {"event": "open", "bulkId": destinatario.message_id, "email": contacto.email},
                   {"event": "open", "bulkId": destinatario.message_id, "email": contacto.email},
                   {"event": "delivery", "bulkId": "msg-hist-cot", "email": contacto.email},
                   {"event": "bounce", "bounceType": "soft_bounce", "bulkId": "msg-hist-aviso", "email": contacto.email},
                   {"event": "delivery", "bulkId": "msg-que-no-existe", "email": "nadie@prueba.ec"}):
        assert client.post(f"/internal/tencent/eventos/{TOKEN_EVENTOS}", json=evento).status_code == 200
    por_tipo = {i["tipo"]: i["estado_entrega"] for i in client.get(ruta, headers=h).json()}
    assert por_tipo == {"BOLETIN": "ABIERTO", "COTIZACION": "ENTREGADO", "AVISO": "REBOTE_SOFT"}
    eventos = db_session.scalars(select(TencentEvento).where(TencentEvento.message_id == destinatario.message_id)).all()
    assert sorted(e.evento for e in eventos) == ["delivery", "open"]  # el repetido no se guarda dos veces
    sin_vincular = db_session.scalars(select(TencentEvento).where(TencentEvento.message_id == "msg-que-no-existe")).one()
    assert (sin_vincular.origen, sin_vincular.origen_id) == ("SIN_VINCULAR", None)
    detalle = client.get(f"/v1/boletines/campanas/{campana['id']}", headers=h).json()
    assert (detalle["entregados"], detalle["abiertos"], detalle["clics"]) == (1, 1, 0)
    previa = client.get(f"{ruta}/BOLETIN/{destinatario.id}", headers=h)
    assert previa.status_code == 200, previa.text
    variables = previa.json()["variables"]
    assert (variables["nombre_contacto"], variables["comercial"]) == (contacto.nombre, esc.cliente.comercial_asignado.nombre)
    assert "/b/" in variables["enlace_baja"] and previa.json()["estado_entrega"] == "Abierto"
    eventos_api = client.get("/v1/notificaciones/eventos-tencent?origen=SIN_VINCULAR", headers=headers_de(crear_usuario("admin")))
    assert eventos_api.status_code == 200 and any(e["message_id"] == "msg-que-no-existe" for e in eventos_api.json())


def test_historial_respeta_el_alcance_del_comercial(client, crear_usuario, headers_de, db_session):  # TE1B-20
    cliente = crear_cliente(db_session, crear_usuario("comercial"))
    ajeno = headers_de(crear_usuario("comercial"))
    assert client.get(f"/v1/clientes/{cliente.id}/correos", headers=ajeno).status_code == 404


def test_columnas_se_guardan_por_usuario(client, crear_usuario, headers_de):  # TE1B-21
    uno, otro = headers_de(crear_usuario("comercial")), headers_de(crear_usuario("comercial"))
    assert client.get("/v1/preferencias/clientes.columnas", headers=uno).json()["valor"] is None
    columnas = ["razon_social", "identificacion", "sector", "ultimo_envio"]
    assert client.put("/v1/preferencias/clientes.columnas", json={"valor": columnas}, headers=uno).status_code == 200
    assert client.get("/v1/preferencias/clientes.columnas", headers=uno).json()["valor"] == columnas
    assert client.get("/v1/preferencias/clientes.columnas", headers=otro).json()["valor"] is None
    assert client.put("/v1/preferencias/clientes.columnas", json={"valor": "texto"}, headers=uno).status_code == 422
    assert client.get("/v1/preferencias/otra.cosa", headers=uno).status_code == 404


def test_exportar_e_importar_da_cero_cambios(client, base, db_session):  # TE1B-22
    email = base["comercial"].email
    a, b = cedula_libre(db_session), cedula_libre(db_session)
    clientes = [_fila_cliente(email, a, etiquetas="vip", volumen_estimado="MEDIO", frecuencia="MENSUAL", notas="Nota  larga"),
                _fila_cliente(email, b, acepta_promociones="NO", fecha_consentimiento="", medio_consentimiento="")]
    contactos = [_fila_contacto(a, f"e1.{a}@prueba.ec", principal="SI", telefono="0991111111"),
                 _fila_contacto(b, f"e2.{b}@prueba.ec", recibe_promociones="NO", area="")]
    assert _confirmar(client, base["h"], "ida.xlsx", _xlsx(clientes, contactos)).status_code == 200
    exportado = client.get("/v1/clientes/exportar.xlsx?tipo=IMPORTADOR", headers=base["h"])
    assert exportado.status_code == 200 and exportado.headers["content-type"].startswith("application/vnd.openxmlformats")
    libro = load_workbook(io.BytesIO(exportado.content))
    identificaciones = [f[1] for f in libro["Clientes"].iter_rows(min_row=3, values_only=True)]
    assert {a, b} <= set(identificaciones)
    vuelta = _previa(client, base["h"], "vuelta.xlsx", exportado.content)
    assert vuelta.status_code == 200, vuelta.text
    resumen = vuelta.json()["resumen"]
    assert (resumen["cambios"], resumen["clientes_error"], resumen["contactos_error"]) == (0, 0, 0)
    assert resumen["clientes_sin_cambios"] >= 2 and resumen["contactos_sin_cambios"] >= 2


def test_exportar_respeta_el_alcance(client, crear_usuario, headers_de, db_session):  # TE1B-23
    propio, ajeno = crear_usuario("comercial"), crear_usuario("comercial")
    mio, otro = crear_cliente(db_session, propio), crear_cliente(db_session, ajeno)
    libro = load_workbook(io.BytesIO(client.get("/v1/clientes/exportar.xlsx", headers=headers_de(propio)).content))
    identificaciones = {f[1] for f in libro["Clientes"].iter_rows(min_row=3, values_only=True)}
    assert mio.ruc_o_id in identificaciones and otro.ruc_o_id not in identificaciones


# ---------- Pantallas (TE1B-24) ----------


@pytest.mark.parametrize("pagina", ["clientes", "listas", "catalogos-perfil", "campana"])
def test_pantallas_sin_estilos_en_linea_ni_html_inyectado(client, pagina):  # TE1B-24
    import re

    from test_frontend import PATRONES_PROHIBIDOS_JS

    respuesta = client.get(f"/app/{pagina}.html")
    assert respuesta.status_code == 200
    assert re.findall(r'<link rel="stylesheet" href="(?!https://)([^"]+)"', respuesta.text) == ["css/xps-pricing.css"]
    assert "<style" not in respuesta.text and "style=" not in respuesta.text
    for script in (f"{pagina}.js", "importacion.js"):
        texto = client.get(f"/app/js/{script}").text
        assert [p for p in PATRONES_PROHIBIDOS_JS if p in texto] == []
    menu = client.get("/app/js/layout.js").text
    assert '"catalogos-perfil.html", permisos: ["config.editar"]' in menu
