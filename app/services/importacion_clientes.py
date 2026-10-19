"""E1-B Importación y exportación de clientes y contactos con la plantilla oficial (docs/extensiones/importacion/).

Columnas EXACTAS de la plantilla aprobada por el gerente:
- Hoja Clientes: COLUMNAS_CLIENTE. Hoja Contactos: COLUMNAS_CONTACTO.
- CSV: una fila por CONTACTO, con las columnas de Clientes y las de Contactos con prefijo "contacto_" (sin
  identificacion_cliente). Separador ; o , (se detecta) y UTF-8 con o sin BOM.

Flujo: leer → validar (errores por fila y columna, en lenguaje simple) → vista previa (crear, actualizar, sin cambios,
error) → confirmar ("solo crear nuevos" o "crear y actualizar") en una sola transacción, con auditoría de cada cliente y
contacto y de la importación (quién, archivo, totales). Nunca borra: una celda opcional vacía no borra el dato guardado.
Las filas de EJEMPLO de la plantilla se ignoran solas. Máximo de filas: config_sistema.importacion_max_filas.

"Descargar plantilla" genera el .xlsx en el momento, sobre la plantilla oficial, con las listas desplegables de los
catálogos vigentes (hoja Catalogos). "Exportar" escribe los clientes en las mismas columnas, para corregirlos en Excel y
volver a importarlos con "crear y actualizar" (sin cambios reales: "0 cambios").
"""
import base64
import csv
import io
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import func, select, update
from sqlalchemy.orm import Session, selectinload

from app.core.errors import ErrorNegocio, confirmar, enviar
from app.core.rbac import permisos_de
from app.models import (
    Cliente,
    ConfigSistema,
    ContactoCliente,
    ListaEnvio,
    Pais,
    Servicio,
    TipoIdentificacion,
    Usuario,
)
from app.schemas.usuarios import email_valido
from app.services import auditoria, perfil_comercial
from app.services.perfil_comercial import clave, lista_de

PLANTILLA = Path(__file__).resolve().parent.parent.parent / "docs" / "extensiones" / "importacion" / "plantilla_importacion_clientes.xlsx"
COLUMNAS_CLIENTE = (
    "tipo_identificacion", "identificacion", "razon_social", "nombre_comercial", "pais", "ciudad", "direccion", "telefono",
    "comercial_email", "tipo_cliente", "estado_comercial", "sector", "servicios_interes", "rutas_interes", "tipos_carga",
    "volumen_estimado", "frecuencia", "origen_cliente", "etiquetas", "acepta_promociones", "fecha_consentimiento",
    "medio_consentimiento", "notas",
)
COLUMNAS_CONTACTO = ("identificacion_cliente", "nombre", "cargo", "area", "email", "telefono", "principal",
                     "recibe_promociones", "recibe_cotizaciones")
COLUMNAS_CSV = (*COLUMNAS_CLIENTE, *[f"contacto_{c}" for c in COLUMNAS_CONTACTO[1:]])
OBLIGATORIAS_CLIENTE = ("tipo_identificacion", "identificacion", "razon_social", "pais", "comercial_email",
                        "acepta_promociones")
OBLIGATORIAS_CONTACTO = ("identificacion_cliente", "nombre", "email")
# Las filas de ejemplo de la plantilla oficial: se ignoran si siguen en el archivo.
EJEMPLOS_CLIENTE = {("0992345678001", "IMPORTADORAEJEMPLOSA"), ("0912345678", "JUANPEREZ")}
EJEMPLOS_CONTACTO = {("0992345678001", "maria.perez@ejemplo.com"), ("0912345678", "juan.perez@ejemplo.com")}
SI = {"SI", "S", "TRUE", "VERDADERO", "1", "YES", "X"}
NO = {"NO", "N", "FALSE", "FALSO", "0"}
MODOS = ("SOLO_NUEVOS", "CREAR_Y_ACTUALIZAR")
TAMANO_MAXIMO = 8 * 1024 * 1024
# Columnas de la hoja Catalogos (las validaciones de la plantilla apuntan a ellas).
COLUMNAS_CATALOGOS = ("tipo_identificacion", "si_no", "TIPO_CLIENTE", "ESTADO_COMERCIAL", "SECTOR", "FRECUENCIA", "VOLUMEN",
                      "ORIGEN_CLIENTE", "AREA_CONTACTO", "MEDIO_CONSENTIMIENTO", "servicios", "TIPO_CARGA")


# ---------- Lectura ----------


def _encabezado(texto: Any) -> str:
    sin_tildes = unicodedata.normalize("NFD", str(texto or "")).encode("ascii", "ignore").decode()
    return "_".join(sin_tildes.strip().lower().replace("*", "").split())


def _texto(valor: Any) -> str:
    """El valor de una celda como texto (los números enteros sin .0, las fechas en dd/mm/aaaa)."""
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return " ".join(str(valor).split())


def _es_descripcion(fila: dict[str, str]) -> bool:
    """La fila 2 de la plantilla oficial explica cada columna ("OBLIGATORIO · …"): no son datos."""
    return any(v.upper().startswith("OBLIGATORIO") for v in fila.values())


def _filas_hoja(hoja, columnas: tuple[str, ...]) -> tuple[list[tuple[int, dict[str, str]]], list[str]]:
    filas = list(hoja.iter_rows(values_only=True))
    if not filas:
        return [], list(columnas)
    encabezados = [_encabezado(c) for c in filas[0]]
    faltan = [c for c in columnas if c not in encabezados]
    salida = []
    for numero, fila in enumerate(filas[1:], start=2):
        datos = {encabezados[i]: _texto(v) for i, v in enumerate(fila) if i < len(encabezados) and encabezados[i]}
        if not any(datos.values()) or (numero == 2 and _es_descripcion(datos)):
            continue
        salida.append((numero, datos))
    return salida, faltan


def leer(nombre: str, contenido: bytes) -> dict[str, Any]:
    """{clientes: [(fila, datos)], contactos: [(fila, datos)], hoja_contactos: str, faltan: [...], formato}."""
    if len(contenido) > TAMANO_MAXIMO:
        raise ErrorNegocio(413, "ARCHIVO_GRANDE", "El archivo pasa de 8 MB")
    extension = Path(nombre or "").suffix.lower()
    if extension == ".xlsx":
        try:
            libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        except Exception:  # noqa: BLE001 - un archivo dañado o que no es Excel
            raise ErrorNegocio(422, "ARCHIVO_INVALIDO", "No se pudo leer el Excel: ¿es la plantilla oficial (.xlsx)?") from None
        hojas = {_encabezado(h): h for h in libro.sheetnames}
        if "clientes" not in hojas or "contactos" not in hojas:
            raise ErrorNegocio(422, "ARCHIVO_INVALIDO", "El Excel debe tener las hojas Clientes y Contactos de la plantilla oficial")
        clientes, faltan_c = _filas_hoja(libro[hojas["clientes"]], COLUMNAS_CLIENTE)
        contactos, faltan_k = _filas_hoja(libro[hojas["contactos"]], COLUMNAS_CONTACTO)
        libro.close()
        return {"clientes": clientes, "contactos": contactos, "faltan": faltan_c + faltan_k, "formato": "xlsx",
                "hoja_clientes": "Clientes", "hoja_contactos": "Contactos"}
    if extension == ".csv":
        texto = None
        for codificacion in ("utf-8-sig", "cp1252"):
            try:
                texto = contenido.decode(codificacion)
                break
            except UnicodeDecodeError:
                continue
        primera = texto.splitlines()[0] if texto and texto.splitlines() else ""
        separador = ";" if primera.count(";") >= primera.count(",") else ","
        lector = csv.reader(io.StringIO(texto), delimiter=separador)
        filas = list(lector)
        if not filas:
            raise ErrorNegocio(422, "ARCHIVO_VACIO", "El archivo está vacío")
        encabezados = [_encabezado(c) for c in filas[0]]
        faltan = [c for c in COLUMNAS_CSV if c not in encabezados]
        clientes: dict[str, tuple[int, dict[str, str]]] = {}
        contactos = []
        for numero, fila in enumerate(filas[1:], start=2):
            datos = {encabezados[i]: _texto(v) for i, v in enumerate(fila) if i < len(encabezados) and encabezados[i]}
            if not any(datos.values()):
                continue
            cliente = {c: datos.get(c, "") for c in COLUMNAS_CLIENTE}
            llave = clave(cliente["tipo_identificacion"]) + "|" + clave(cliente["identificacion"])
            if llave not in clientes:
                clientes[llave] = (numero, cliente)
            contacto = {c: datos.get(f"contacto_{c}", "") for c in COLUMNAS_CONTACTO[1:]}
            if any(contacto.values()):
                contactos.append((numero, {"identificacion_cliente": cliente["identificacion"]} | contacto))
        return {"clientes": list(clientes.values()), "contactos": contactos, "faltan": faltan, "formato": "csv",
                "hoja_clientes": "CSV", "hoja_contactos": "CSV"}
    raise ErrorNegocio(422, "ARCHIVO_INVALIDO", "Suba la plantilla oficial en .xlsx o .csv")


def decodificar(contenido_base64: str) -> bytes:
    try:
        return base64.b64decode(contenido_base64.split(",")[-1], validate=False)
    except (ValueError, TypeError):
        raise ErrorNegocio(422, "ARCHIVO_INVALIDO", "No se pudo leer el archivo enviado") from None


# ---------- Validación ----------


def _si_no(valor: str, *, por_defecto: bool | None = None) -> bool | None:
    if not valor:
        return por_defecto
    llave = clave(valor)
    if llave in SI:
        return True
    if llave in NO:
        return False
    raise ValueError(valor)


def _fecha(valor: str) -> date | None:
    if not valor:
        return None
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(valor, formato).date()
        except ValueError:
            continue
    raise ValueError(valor)


class Contexto:
    """Lo que se consulta una sola vez por importación."""

    def __init__(self, db: Session, actor: Usuario):
        self.db, self.actor = db, actor
        self.cache: dict[str, Any] = {}
        self.ve_todos = "cotizacion.ver_todas" in permisos_de(actor)
        self.tipos_id = {}
        for tipo in db.scalars(select(TipoIdentificacion).where(TipoIdentificacion.activo.is_(True))):
            self.tipos_id[clave(tipo.codigo)] = tipo.codigo
            self.tipos_id.setdefault(clave(tipo.nombre), tipo.codigo)
        self.paises = {p.iso2: p for p in db.scalars(select(Pais).where(Pais.activo.is_(True)))}
        self.usuarios = {u.email.lower(): u for u in db.scalars(select(Usuario).where(Usuario.activo.is_(True)))}


def _validar_cliente(ctx: Contexto, datos: dict[str, str]) -> tuple[dict[str, Any], list[dict[str, str]]]:
    from app.services import clientes as servicio_clientes

    errores: list[dict[str, str]] = []

    def error(columna: str, mensaje: str) -> None:
        errores.append({"columna": columna, "mensaje": mensaje})

    for columna in OBLIGATORIAS_CLIENTE:
        if not datos.get(columna):
            error(columna, "Es obligatorio")
    valores: dict[str, Any] = {}
    tipo = ctx.tipos_id.get(clave(datos.get("tipo_identificacion")))
    if datos.get("tipo_identificacion") and tipo is None:
        error("tipo_identificacion", f"«{datos['tipo_identificacion']}» no es un tipo de identificación del catálogo")
    valores["tipo_id"] = tipo
    if tipo and datos.get("identificacion"):
        try:
            valores["ruc_o_id"] = servicio_clientes.normalizar_identificacion(ctx.db, tipo, datos["identificacion"])
        except ErrorNegocio as fallo:
            error("identificacion", fallo.mensaje)
    valores["razon_social"] = " ".join(datos.get("razon_social", "").split()) or None
    pais = ctx.paises.get((datos.get("pais") or "").strip().upper())
    if datos.get("pais") and pais is None:
        error("pais", f"«{datos['pais']}» no es un país activo (use el código ISO de 2 letras, por ejemplo EC)")
    valores["pais_id"] = pais.id if pais else None
    usuario = ctx.usuarios.get((datos.get("comercial_email") or "").strip().lower())
    if datos.get("comercial_email"):
        if usuario is None or "cotizacion.crear" not in permisos_de(usuario):
            error("comercial_email", f"«{datos['comercial_email']}» no es un usuario activo que pueda cotizar")
        elif not ctx.ve_todos and usuario.id != ctx.actor.id:
            error("comercial_email", "Solo puede importar clientes asignados a usted")
    valores["comercial_asignado_id"] = usuario.id if usuario else None
    for campo in ("ciudad", "direccion", "telefono", "notas"):
        valores[campo] = " ".join(datos.get(campo, "").split()) or None
    perfil_entrada = {c: datos.get(c, "") for c in ("nombre_comercial", *perfil_comercial.CAMPOS_CATALOGO,
                                                       *perfil_comercial.CAMPOS_LISTA)}
    try:
        perfil_entrada["acepta_promociones"] = _si_no(datos.get("acepta_promociones", ""))
    except ValueError:
        error("acepta_promociones", "Use SI o NO")
    try:
        perfil_entrada["fecha_consentimiento"] = _fecha(datos.get("fecha_consentimiento", ""))
    except ValueError:
        error("fecha_consentimiento", "Use el formato dd/mm/aaaa")
    perfil, fallos = perfil_comercial.validar_perfil(ctx.db, perfil_entrada, ctx.cache)
    for campo, mensaje in fallos:
        error(campo, mensaje)
    valores |= perfil
    return valores, errores


def _validar_contacto(ctx: Contexto, datos: dict[str, str]) -> tuple[dict[str, Any], list[dict[str, str]]]:
    errores: list[dict[str, str]] = []

    def error(columna: str, mensaje: str) -> None:
        errores.append({"columna": columna, "mensaje": mensaje})

    for columna in OBLIGATORIAS_CONTACTO:
        if not datos.get(columna):
            error(columna, "Es obligatorio")
    valores: dict[str, Any] = {"nombre": " ".join(datos.get("nombre", "").split()) or None}
    email = (datos.get("email") or "").strip().lower()
    if email and not email_valido(email):
        error("email", f"«{datos['email']}» no es un correo válido")
    valores["email"] = email or None
    for campo in ("cargo", "telefono"):
        valores[campo] = " ".join(datos.get(campo, "").split()) or None
    if datos.get("area"):
        area = perfil_comercial.resolver(ctx.db, "AREA_CONTACTO", datos["area"], ctx.cache)
        if area is None:
            error("area", f"«{datos['area']}» no está en el catálogo de áreas")
        valores["area"] = area
    for columna, campo, defecto in (("principal", "principal", False), ("recibe_promociones", "recibe_promociones", True),
                                    ("recibe_cotizaciones", "recibir_avisos", True)):
        try:
            valores[campo] = _si_no(datos.get(columna, ""), por_defecto=defecto)
        except ValueError:
            error(columna, "Use SI o NO")
    return valores, errores


def _cambios(objeto: Any, valores: dict[str, Any], campos: tuple[str, ...]) -> dict[str, Any]:
    """Los campos que cambiarían. Un valor vacío del archivo no borra el guardado (nunca se borra)."""
    salida = {}
    for campo in campos:
        nuevo = valores.get(campo)
        if nuevo in (None, "", []):
            continue
        actual = getattr(objeto, campo)
        if isinstance(nuevo, list):
            iguales = list(actual or []) == nuevo
        elif isinstance(nuevo, str):
            iguales = " ".join(str(actual or "").split()) == " ".join(nuevo.split())
        else:
            iguales = actual == nuevo
        if not iguales:
            salida[campo] = nuevo
    return salida


CAMPOS_CLIENTE_ACTUALIZABLES = ("razon_social", "pais_id", "comercial_asignado_id", "ciudad", "direccion", "telefono", "notas",
                                *perfil_comercial.CAMPOS_PERFIL)
CAMPOS_CONTACTO_ACTUALIZABLES = ("nombre", "cargo", "area", "telefono", "principal", "recibe_promociones", "recibir_avisos")


def analizar(db: Session, actor: Usuario, nombre: str, contenido: bytes) -> dict[str, Any]:
    """Valida todo el archivo sin guardar nada. Es la base de la vista previa y de la confirmación."""
    leido = leer(nombre, contenido)
    if leido["faltan"]:
        raise ErrorNegocio(422, "COLUMNAS_FALTANTES", "Faltan columnas de la plantilla oficial: " + ", ".join(leido["faltan"]),
                           extra={"faltan": leido["faltan"]})
    maximo = db.get(ConfigSistema, 1).importacion_max_filas
    if len(leido["clientes"]) + len(leido["contactos"]) > maximo * 2 or max(len(leido["clientes"]), len(leido["contactos"])) > maximo:
        raise ErrorNegocio(422, "DEMASIADAS_FILAS", f"El archivo pasa del máximo de {maximo} filas por importación")
    ctx = Contexto(db, actor)
    ejemplos_ignorados = 0
    clientes: list[dict[str, Any]] = []
    por_identificacion: dict[str, dict[str, Any]] = {}
    for numero, datos in leido["clientes"]:
        if (clave(datos.get("identificacion")), clave(datos.get("razon_social"))) in {(a, b) for a, b in EJEMPLOS_CLIENTE}:
            ejemplos_ignorados += 1
            continue
        valores, errores = _validar_cliente(ctx, datos)
        fila = {"hoja": leido["hoja_clientes"], "fila": numero, "tipo": "cliente", "identificacion": datos.get("identificacion", ""),
                "nombre": datos.get("razon_social", ""), "valores": valores, "errores": errores}
        llave = clave(valores.get("ruc_o_id") or datos.get("identificacion"))
        if llave and llave in por_identificacion:
            errores.append({"columna": "identificacion", "mensaje": f"Repetida en el archivo (ya está en la fila {por_identificacion[llave]['fila']})"})
        elif llave:
            por_identificacion[llave] = fila
        existente = None
        if not errores and valores.get("tipo_id") and valores.get("ruc_o_id"):
            existente = db.scalars(select(Cliente).where(Cliente.tipo_id == valores["tipo_id"],
                                                         Cliente.ruc_o_id == valores["ruc_o_id"])).one_or_none()
            if existente is not None and not ctx.ve_todos and existente.comercial_asignado_id != actor.id:
                errores.append({"columna": "identificacion", "mensaje": "Ese cliente ya existe y está asignado a otro comercial"})
        fila["existente"] = existente
        if errores:
            fila["accion"] = "ERROR"
        elif existente is None:
            fila["accion"] = "CREAR"
        else:
            fila["cambios"] = _cambios(existente, valores, CAMPOS_CLIENTE_ACTUALIZABLES)
            fila["accion"] = "ACTUALIZAR" if fila["cambios"] else "SIN_CAMBIOS"
        clientes.append(fila)

    contactos: list[dict[str, Any]] = []
    principales: dict[str, int] = {}
    emails_por_cliente: dict[tuple[str, str], int] = {}
    for numero, datos in leido["contactos"]:
        if (clave(datos.get("identificacion_cliente")), (datos.get("email") or "").strip().lower()) in EJEMPLOS_CONTACTO:
            ejemplos_ignorados += 1
            continue
        valores, errores = _validar_contacto(ctx, datos)
        fila = {"hoja": leido["hoja_contactos"], "fila": numero, "tipo": "contacto",
                "identificacion": datos.get("identificacion_cliente", ""), "nombre": datos.get("nombre", ""),
                "email": valores.get("email") or "", "valores": valores, "errores": errores}
        llave = clave(datos.get("identificacion_cliente"))
        padre = por_identificacion.get(llave)
        cliente_bd = None
        if padre is None and llave:
            cliente_bd = db.scalars(select(Cliente).where(Cliente.ruc_o_id == llave.upper())).first()
            if cliente_bd is None:
                errores.append({"columna": "identificacion_cliente",
                                "mensaje": "No hay un cliente con esa identificación en la hoja Clientes ni en el sistema"})
            elif not ctx.ve_todos and cliente_bd.comercial_asignado_id != actor.id:
                errores.append({"columna": "identificacion_cliente", "mensaje": "Ese cliente está asignado a otro comercial"})
        elif padre is not None and padre["accion"] == "ERROR":
            errores.append({"columna": "identificacion_cliente", "mensaje": f"Su cliente tiene errores (fila {padre['fila']})"})
        if valores.get("principal") and llave:
            if llave in principales:
                errores.append({"columna": "principal", "mensaje": f"El cliente ya tiene otro contacto principal en el archivo (fila {principales[llave]})"})
            else:
                principales[llave] = numero
        if valores.get("email") and llave:
            repetido = emails_por_cliente.get((llave, valores["email"]))
            if repetido:
                errores.append({"columna": "email", "mensaje": f"Repetido para el mismo cliente (fila {repetido})"})
            else:
                emails_por_cliente[(llave, valores["email"])] = numero
        cliente_existente = padre["existente"] if padre is not None else cliente_bd
        existente = None
        if cliente_existente is not None and valores.get("email"):
            existente = db.scalars(select(ContactoCliente).where(
                ContactoCliente.cliente_id == cliente_existente.id,
                func.lower(ContactoCliente.email) == valores["email"])).first()
        fila["padre"], fila["cliente_bd"], fila["existente"] = padre, cliente_existente, existente
        if errores:
            fila["accion"] = "ERROR"
        elif existente is None:
            fila["accion"] = "CREAR"
        else:
            fila["cambios"] = _cambios(existente, valores, CAMPOS_CONTACTO_ACTUALIZABLES)
            if existente.principal and valores.get("principal") is False and "principal" in fila["cambios"]:
                fila["cambios"].pop("principal")  # un archivo sin "principal" no le quita el principal a nadie
            fila["accion"] = "ACTUALIZAR" if fila["cambios"] else "SIN_CAMBIOS"
        contactos.append(fila)
    return {"clientes": clientes, "contactos": contactos, "ejemplos_ignorados": ejemplos_ignorados,
            "formato": leido["formato"]}


def _resumen(analisis: dict[str, Any]) -> dict[str, int]:
    contar = lambda filas, accion: sum(1 for f in filas if f["accion"] == accion)  # noqa: E731
    c, k = analisis["clientes"], analisis["contactos"]
    return {
        "clientes_nuevos": contar(c, "CREAR"), "clientes_actualizar": contar(c, "ACTUALIZAR"),
        "clientes_sin_cambios": contar(c, "SIN_CAMBIOS"), "clientes_error": contar(c, "ERROR"),
        "contactos_nuevos": contar(k, "CREAR"), "contactos_actualizar": contar(k, "ACTUALIZAR"),
        "contactos_sin_cambios": contar(k, "SIN_CAMBIOS"), "contactos_error": contar(k, "ERROR"),
        "ejemplos_ignorados": analisis["ejemplos_ignorados"],
        "cambios": contar(c, "CREAR") + contar(c, "ACTUALIZAR") + contar(k, "CREAR") + contar(k, "ACTUALIZAR"),
    }


def _fila_publica(fila: dict[str, Any], resultado: str | None = None) -> dict[str, Any]:
    return {"hoja": fila["hoja"], "fila": fila["fila"], "tipo": fila["tipo"], "identificacion": fila["identificacion"],
            "nombre": fila["nombre"], "email": fila.get("email", ""), "accion": fila["accion"],
            "cambios": sorted(fila.get("cambios", {})), "errores": fila["errores"], "resultado": resultado or fila["accion"]}


def vista_previa(db: Session, actor: Usuario, nombre: str, contenido: bytes) -> dict[str, Any]:
    analisis = analizar(db, actor, nombre, contenido)
    db.rollback()
    return {"resumen": _resumen(analisis),
            "filas": [_fila_publica(f) for f in analisis["clientes"] + analisis["contactos"]]}


def _foto(objeto: Any, campos: tuple[str, ...]) -> dict[str, Any]:
    salida = {}
    for campo in campos:
        valor = getattr(objeto, campo)
        salida[campo] = valor.isoformat() if isinstance(valor, date) else valor
    return salida


def confirmar_importacion(db: Session, actor: Usuario, nombre: str, contenido: bytes, modo: str,
                          lista_id: int | None = None) -> dict[str, Any]:
    if modo not in MODOS:
        raise ErrorNegocio(422, "MODO_INVALIDO", "Elija «solo crear nuevos» o «crear y actualizar»")
    lista = None
    if lista_id is not None:
        lista = db.get(ListaEnvio, lista_id)
        if lista is None or lista.tipo != "MANUAL":
            raise ErrorNegocio(422, "LISTA_INVALIDA", "Los contactos importados solo se agregan a una lista manual")
    analisis = analizar(db, actor, nombre, contenido)
    actualizar = modo == "CREAR_Y_ACTUALIZAR"
    resultados: list[dict[str, Any]] = []
    tocados: list[ContactoCliente] = []
    nuevos_por_fila: dict[int, Cliente] = {}
    totales = {"clientes_creados": 0, "clientes_actualizados": 0, "contactos_creados": 0, "contactos_actualizados": 0,
               "omitidos": 0, "errores": 0}
    campos_foto_cliente = ("razon_social", "tipo_id", "ruc_o_id", "pais_id", "comercial_asignado_id", "ciudad", "direccion",
                           "telefono", "notas", *perfil_comercial.CAMPOS_PERFIL)
    for fila in analisis["clientes"]:
        valores = fila["valores"]
        if fila["accion"] == "ERROR":
            totales["errores"] += 1
            resultados.append(_fila_publica(fila, "ERROR"))
            continue
        if fila["accion"] == "CREAR":
            cliente = Cliente(
                razon_social=valores["razon_social"], tipo_id=valores["tipo_id"], ruc_o_id=valores["ruc_o_id"],
                pais_id=valores["pais_id"], comercial_asignado_id=valores["comercial_asignado_id"],
                ciudad=valores.get("ciudad"), direccion=valores.get("direccion"), telefono=valores.get("telefono"),
                notas=valores.get("notas"), creado_por=auditoria.id_de(actor),
                **{c: valores[c] for c in perfil_comercial.CAMPOS_PERFIL if c in valores and valores[c] not in (None, [])},
            )
            if cliente.estado_comercial is None:
                cliente.estado_comercial = "ACTIVO"
            db.add(cliente)
            enviar(db)
            auditoria.registrar(db, actor, "cliente", cliente.id, "CREAR", despues=_foto(cliente, campos_foto_cliente),
                                nota=f"IMPORTAR {nombre}"[:300])
            nuevos_por_fila[fila["fila"]] = cliente
            totales["clientes_creados"] += 1
            resultados.append(_fila_publica(fila, "CREADO"))
        elif fila["accion"] == "ACTUALIZAR" and actualizar:
            cliente = fila["existente"]
            antes = _foto(cliente, campos_foto_cliente)
            for campo, valor in fila["cambios"].items():
                setattr(cliente, campo, valor)
            auditoria.registrar(db, actor, "cliente", cliente.id, "EDITAR", antes=antes,
                                despues=_foto(cliente, campos_foto_cliente), nota=f"IMPORTAR {nombre}"[:300])
            totales["clientes_actualizados"] += 1
            resultados.append(_fila_publica(fila, "ACTUALIZADO"))
        else:
            totales["omitidos"] += 1
            resultados.append(_fila_publica(fila, "SIN_CAMBIOS" if fila["accion"] == "SIN_CAMBIOS" else "OMITIDO (solo crear nuevos)"))
    enviar(db)
    campos_foto_contacto = ("nombre", "cargo", "area", "email", "telefono", "principal", "recibe_promociones", "recibir_avisos")
    for fila in analisis["contactos"]:
        valores = fila["valores"]
        if fila["accion"] == "ERROR":
            totales["errores"] += 1
            resultados.append(_fila_publica(fila, "ERROR"))
            continue
        cliente = fila["cliente_bd"] or (nuevos_por_fila.get(fila["padre"]["fila"]) if fila["padre"] else None)
        if cliente is None:
            totales["omitidos"] += 1
            resultados.append(_fila_publica(fila, "OMITIDO (su cliente no se importó)"))
            continue
        if valores.get("principal"):
            db.execute(update(ContactoCliente).where(ContactoCliente.cliente_id == cliente.id,
                                                     ContactoCliente.principal.is_(True))
                       .values(principal=False).execution_options(synchronize_session="fetch"))
            enviar(db)
        if fila["accion"] == "CREAR":
            contacto = ContactoCliente(
                cliente_id=cliente.id, nombre=valores["nombre"], email=valores["email"], cargo=valores.get("cargo"),
                telefono=valores.get("telefono"), area=valores.get("area"), principal=bool(valores.get("principal")),
                recibe_promociones=valores.get("recibe_promociones", True),
                recibir_avisos=valores.get("recibir_avisos", True), creado_por=auditoria.id_de(actor),
            )
            db.add(contacto)
            enviar(db)
            auditoria.registrar(db, actor, "contacto_cliente", contacto.id, "CREAR",
                                despues=_foto(contacto, campos_foto_contacto), nota=f"IMPORTAR {nombre}"[:300])
            totales["contactos_creados"] += 1
            tocados.append(contacto)
            resultados.append(_fila_publica(fila, "CREADO"))
        elif fila["accion"] == "ACTUALIZAR" and actualizar:
            contacto = fila["existente"]
            antes = _foto(contacto, campos_foto_contacto)
            for campo, valor in fila["cambios"].items():
                setattr(contacto, campo, valor)
            enviar(db)
            auditoria.registrar(db, actor, "contacto_cliente", contacto.id, "EDITAR", antes=antes,
                                despues=_foto(contacto, campos_foto_contacto), nota=f"IMPORTAR {nombre}"[:300])
            totales["contactos_actualizados"] += 1
            tocados.append(contacto)
            resultados.append(_fila_publica(fila, "ACTUALIZADO"))
        else:
            if fila["existente"] is not None:
                tocados.append(fila["existente"])
            totales["omitidos"] += 1
            resultados.append(_fila_publica(fila, "SIN_CAMBIOS" if fila["accion"] == "SIN_CAMBIOS" else "OMITIDO (solo crear nuevos)"))
    agregados = 0
    if lista is not None:
        actuales = {c.id for c in lista.miembros}
        nuevos = [c for c in tocados if c.id not in actuales and c.activo]
        lista.miembros = list(lista.miembros) + nuevos
        agregados = len(nuevos)
    resumen = _resumen(analisis) | totales | {"agregados_a_lista": agregados}
    auditoria.registrar(db, actor, "importacion_clientes", nombre[:40] or "archivo", "CREAR",
                        despues={"archivo": nombre, "modo": modo, "lista_id": lista_id} | totales,
                        nota=f"IMPORTAR {nombre}"[:300])
    confirmar(db)
    return {"resumen": resumen, "filas": resultados}


# ---------- Plantilla, CSV de ejemplo y exportación ----------


def _valores_catalogos(db: Session) -> dict[str, list[str]]:
    salida = {
        "tipo_identificacion": [t.codigo for t in db.scalars(select(TipoIdentificacion).where(
            TipoIdentificacion.activo.is_(True)).order_by(TipoIdentificacion.orden, TipoIdentificacion.codigo))],
        "si_no": ["SI", "NO"],
        "servicios": [s.codigo for s in db.scalars(select(Servicio).where(Servicio.activo.is_(True)).order_by(Servicio.orden, Servicio.codigo))],
    }
    for tipo in ("TIPO_CLIENTE", "ESTADO_COMERCIAL", "SECTOR", "FRECUENCIA", "VOLUMEN", "ORIGEN_CLIENTE", "AREA_CONTACTO",
                 "MEDIO_CONSENTIMIENTO", "TIPO_CARGA"):
        salida[tipo] = [(v.nombre if tipo == "VOLUMEN" else v.codigo) for v in perfil_comercial.valores(db, tipo)]
    return salida


def _letra(indice: int) -> str:
    return chr(ord("A") + indice)


def _libro_con_catalogos(db: Session):
    """La plantilla oficial con la hoja Catalogos y sus listas desplegables al día con la base."""
    libro = load_workbook(PLANTILLA)
    hoja = libro["Catalogos"]
    valores = _valores_catalogos(db)
    for fila in hoja.iter_rows(min_row=2, max_row=max(hoja.max_row, 2)):
        for celda in fila:
            celda.value = None
    finales = {}
    for indice, columna in enumerate(COLUMNAS_CATALOGOS):
        lista = valores[columna]
        for numero, valor in enumerate(lista, start=2):
            hoja.cell(row=numero, column=indice + 1, value=valor)
        finales[_letra(indice)] = len(lista) + 1
    for nombre_hoja in ("Clientes", "Contactos"):
        destino = libro[nombre_hoja]
        nuevas = []
        for validacion in destino.data_validations.dataValidation:
            formula = validacion.formula1 or ""
            encontrado = re.match(r"=?Catalogos!\$([A-Z])\$2:\$[A-Z]\$\d+", formula)
            if encontrado:
                letra = encontrado.group(1)
                validacion.formula1 = f"=Catalogos!${letra}$2:${letra}${max(finales.get(letra, 2), 2)}"
            nuevas.append(validacion)
        destino.data_validations.dataValidation = nuevas
    return libro


def plantilla_xlsx(db: Session) -> bytes:
    salida = io.BytesIO()
    _libro_con_catalogos(db).save(salida)
    return salida.getvalue()


FILAS_EJEMPLO_CSV = [
    ["RUC", "0992345678001", "IMPORTADORA EJEMPLO S.A.", "Ejemplo Import", "EC", "Guayaquil", "Av. Ejemplo 123", "042000000",
     "comercial@xpress.ec", "IMPORTADOR", "ACTIVO", "TEXTIL", "FCL|LCL", "CN>EC|US>EC", "GENERAL",
     "MEDIO (1 a 5 TON o 1 a 3 contenedores al mes)", "MENSUAL", "REFERIDO", "VIP|FERIA_2026", "SI", "19/10/2026",
     "RELACION_COMERCIAL_PREVIA", "", "María Pérez", "Jefa de Compras", "COMPRAS", "maria.perez@ejemplo.com", "0991111111",
     "SI", "SI", "SI"],
    ["CEDULA", "0912345678", "JUAN PÉREZ", "", "EC", "Quito", "", "0990000000", "comercial@xpress.ec", "PERSONA_NATURAL",
     "PROSPECTO", "E_COMMERCE", "AEREO", "US>EC", "GENERAL|PERECIBLE", "BAJO (menos de 1 TON o 1 contenedor al mes)",
     "OCASIONAL", "WEB", "", "SI", "19/10/2026", "FORMULARIO_WEB", "", "Juan Pérez", "Gerente", "GERENCIA",
     "juan.perez@ejemplo.com", "0990000000", "SI", "SI", "SI"],
]


def plantilla_csv() -> bytes:
    salida = io.StringIO()
    escritor = csv.writer(salida, delimiter=";", lineterminator="\n")
    escritor.writerow(COLUMNAS_CSV)
    escritor.writerows(FILAS_EJEMPLO_CSV)
    return ("﻿" + salida.getvalue()).encode("utf-8")


def _si(valor: bool | None) -> str:
    return "SI" if valor else "NO"


def _nombre_volumen(db: Session, codigo: str | None, cache: dict) -> str:
    if not codigo:
        return ""
    if "volumen" not in cache:
        cache["volumen"] = {v.codigo: v.nombre for v in perfil_comercial.valores(db, "VOLUMEN", solo_activos=False)}
    return cache["volumen"].get(codigo, codigo)


def fila_cliente(db: Session, cliente: Cliente, cache: dict) -> list[str]:
    pais = cliente.pais.iso2 if cliente.pais else ""
    return [
        cliente.tipo_id, cliente.ruc_o_id, cliente.razon_social, cliente.nombre_comercial or "", pais, cliente.ciudad or "",
        cliente.direccion or "", cliente.telefono or "", cliente.comercial_asignado.email if cliente.comercial_asignado else "",
        cliente.tipo_cliente or "", cliente.estado_comercial or "", cliente.sector or "", "|".join(cliente.servicios_interes or []),
        "|".join(cliente.rutas_interes or []), "|".join(cliente.tipos_carga or []),
        _nombre_volumen(db, cliente.volumen_estimado, cache), cliente.frecuencia or "", cliente.origen_cliente or "",
        "|".join(cliente.etiquetas or []), _si(cliente.acepta_promociones),
        cliente.fecha_consentimiento.strftime("%d/%m/%Y") if cliente.fecha_consentimiento else "",
        cliente.medio_consentimiento or "", cliente.notas or "",
    ]


def fila_contacto(cliente: Cliente, contacto: ContactoCliente) -> list[str]:
    return [cliente.ruc_o_id, contacto.nombre, contacto.cargo or "", contacto.area or "", contacto.email or "",
            contacto.telefono or "", _si(contacto.principal), _si(contacto.recibe_promociones), _si(contacto.recibir_avisos)]


def exportar_xlsx(db: Session, clientes: list[Cliente]) -> bytes:
    """Los clientes (y sus contactos activos con correo) en las columnas de la plantilla oficial."""
    libro = _libro_con_catalogos(db)
    cache: dict = {}
    for nombre_hoja, filas in (
        ("Clientes", [fila_cliente(db, c, cache) for c in clientes]),
        ("Contactos", [fila_contacto(c, k) for c in clientes for k in sorted(c.contactos, key=lambda x: x.id)
                       if k.activo and k.email]),
    ):
        hoja = libro[nombre_hoja]
        for fila in hoja.iter_rows(min_row=3, max_row=max(hoja.max_row, 3)):
            for celda in fila:
                celda.value = None
        for numero, fila in enumerate(filas, start=3):
            for columna, valor in enumerate(fila, start=1):
                celda = hoja.cell(row=numero, column=columna, value=valor)
                celda.number_format = "@"  # texto: la identificación conserva sus ceros
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def clientes_para_exportar(db: Session, consulta) -> list[Cliente]:
    return list(db.scalars(consulta.options(selectinload(Cliente.contactos), selectinload(Cliente.pais),
                                            selectinload(Cliente.comercial_asignado))))

